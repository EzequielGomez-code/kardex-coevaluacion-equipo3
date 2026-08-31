"""
Constructor único de reportes en Excel.

Mismo criterio que ``pdf.py``: encabezado con el color de marca, columnas
anchas según su contenido y una sola función para los cuatro reportes — no
un volcado crudo de filas sin formato.
"""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLOR_ACENTO = "1F5B4E"
ANCHO_MAXIMO_COLUMNA = 42


def construir_excel(titulo: str, encabezados: list[str], filas: list[list[str]]) -> bytes:
    """Arma un libro de una sola hoja, con encabezado destacado y columnas ajustadas."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = titulo[:31] or "Reporte"  # Excel no admite más de 31 caracteres.

    hoja.append(encabezados)
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=COLOR_ACENTO)
        celda.alignment = Alignment(horizontal="center", vertical="center")

    for fila in filas:
        hoja.append(fila)

    anchos = [len(str(encabezado)) for encabezado in encabezados]
    for fila in filas:
        for indice, valor in enumerate(fila):
            anchos[indice] = max(anchos[indice], len(str(valor)))
    for indice, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(indice)].width = min(
            ancho + 2, ANCHO_MAXIMO_COLUMNA
        )

    hoja.freeze_panes = "A2"

    buffer = io.BytesIO()
    libro.save(buffer)
    return buffer.getvalue()
